from pathlib import Path
from datetime import datetime
from openpyxl import Workbook, load_workbook
import logging
import os
from email.message import EmailMessage
import smtplib
from dotenv import load_dotenv

# ============================================================
# CONFIGURAÇÃO
# ============================================================

load_dotenv()
EMAIL_REMETENTE = os.getenv("EMAIL_REMETENTE")
EMAIL_SENHA = os.getenv("EMAIL_SENHA")
SMTP_SERVER = os.getenv("SMTP_SERVER", "smtp.gmail.com")
SMTP_PORT = int(os.getenv("SMTP_PORT", "465"))

#print("EMAIL:", EMAIL_REMETENTE)
#print("SENHA APP carregada:", bool(EMAIL_SENHA))

class ErroPlanilhaMestra(Exception):
    """Erro ao acessar ou carregar a planilha mestra."""
    pass


def get_project_root():
    current = Path(__file__).resolve().parent
    for parent in [current] + list(current.parents):
        if (parent / "ERP_Portal_Fake").exists() or (parent / "HyperAutomation").exists():
            return parent
    return Path(__file__).resolve().parents[3]


def get_default_planilha_mestra():
    root = get_project_root()
    p1 = root / "ERP_Portal_Fake" / "Sistema_Integrador_Portal_Fake" / "Planilha_Mestra.xlsx"
    p2 = root / "ERP_Portal_Fake" / "Sistema_Integrador_Portal_Fake" / "planilha_mestra.xlsx"
    if p1.exists():
        return p1
    if p2.exists():
        return p2
    return p1


def get_default_planilha_sac():
    root = get_project_root()
    return root / "ERP_Portal_Fake" / "Sistema_Integrador_Portal_Fake" / "atendimentos_sac.xlsx"


PLANILHA_MESTRA = get_default_planilha_mestra()
PLANILHA_SAC = get_default_planilha_sac()


# ============================================================
# LOG
# ============================================================

logger = logging.getLogger("processo4_sac")
if not logger.handlers:
    logger.setLevel(logging.INFO)
    formatter = logging.Formatter("%(asctime)s | %(levelname)s | %(message)s")
    
    log_dir = get_project_root() / "logs"
    log_dir.mkdir(parents=True, exist_ok=True)
    file_handler = logging.FileHandler(str(log_dir / "processo4_sac.log"), encoding="utf-8")
    file_handler.setFormatter(formatter)
    
    stream_handler = logging.StreamHandler()
    stream_handler.setFormatter(formatter)
    
    logger.addHandler(file_handler)
    logger.addHandler(stream_handler)


# ============================================================
# REGISTRO DE ATENDIMENTO
# ============================================================

class RegistroAtendimento:

    def __init__(self, caminho):
        self.caminho = Path(caminho)

    def criar_planilha(self):

        if self.caminho.exists():
            return

        self.caminho.parent.mkdir(
            parents=True,
            exist_ok=True
        )

        wb = Workbook()
        ws = wb.active
        ws.title = "Atendimentos"

        cabecalhos = [
            "Chave Atendimento",
            "Protocolo",
            "CPF",
            "Nome",
            "E-mail",
            "Status Cadastro",
            "Tipo Atendimento",
            "Status Atendimento",
            "Data Atendimento",
            "Observação"
        ]

        ws.append(cabecalhos)

        wb.save(self.caminho)

        logger.info(
            "Planilha de atendimento criada: %s",
            self.caminho
        )

    def carregar(self):

        self.criar_planilha()

        self.wb = load_workbook(self.caminho)
        self.ws = self.wb["Atendimentos"]

    def ja_processado(self, chave_atendimento):

        for linha in range(2, self.ws.max_row + 1):

            chave = self.ws.cell(
                linha,
                1
            ).value

            if chave == chave_atendimento:
                return True

        return False

    def registrar(
        self,
        cliente,
        tipo_atendimento,
        status_atendimento,
        data_atendimento,
        observacao
    ):

        protocolo = cliente["Protocolo"]

        chave_atendimento = (
            f"{protocolo}_{tipo_atendimento}"
        )

        self.ws.append([
            chave_atendimento,
            protocolo,
            cliente["CPF"],
            cliente["Nome"],
            cliente["E-mail"],
            cliente["Status"],
            tipo_atendimento,
            status_atendimento,
            data_atendimento,
            observacao
        ])

        self.wb.save(self.caminho)

        logger.info(
            "Atendimento registrado: %s",
            chave_atendimento
        )


# ============================================================
# DEFINIR MENSAGEM DA NOTIFICACAO
# ============================================================

def criar_mensagem(cliente):

    nome = cliente["Nome"]
    protocolo = cliente["Protocolo"]

    data_atendimento = datetime.now().strftime(
        "%d/%m/%Y %H:%M"
    )

    assunto = (
        f"Cadastro realizado com sucesso - "
        f"Protocolo {protocolo}"
    )

    mensagem = f"""Olá, {nome}!

Informamos que seu cadastro foi realizado com sucesso.

Data do atendimento: {data_atendimento}
Protocolo: {protocolo}

Guarde este protocolo para futuras consultas.

Atenciosamente,
Equipe de Atendimento
"""

    return assunto, mensagem, data_atendimento

# ============================================================
# ENVIAR E-MAIL
# ============================================================

def enviar_comunicacao(cliente):

    email_destino = cliente.get("E-mail")

    if not email_destino:
        raise ValueError(
            f"Cliente {cliente.get('Nome')} "
            "não possui e-mail."
        )

    assunto, mensagem, data_atendimento = (
        criar_mensagem(cliente)
    )

    email = EmailMessage()

    email["From"] = EMAIL_REMETENTE
    email["To"] = email_destino
    email["Subject"] = assunto

    email.set_content(mensagem)

    logger.info(
        "Enviando comunicação para %s",
        email_destino
    )

    with smtplib.SMTP(SMTP_SERVER, SMTP_PORT) as servidor:
        servidor.starttls()
        servidor.login(EMAIL_REMETENTE, EMAIL_SENHA)
        servidor.send_message(email)

    logger.info(
        "Comunicação enviada com sucesso para %s",
        email_destino
    )

    return data_atendimento


# ============================================================
# PROCESSO SAC
# ============================================================

class ProcessoSAC:

    def __init__(
        self,
        caminho_planilha_mestra=None,
        caminho_planilha_sac=None
    ):

        self.caminho_mestra = Path(
            caminho_planilha_mestra or PLANILHA_MESTRA
        )

        self.registro = RegistroAtendimento(
            caminho_planilha_sac or PLANILHA_SAC
        )

    def carregar_planilha_mestra(self):

        try:
            if not self.caminho_mestra.exists():
                alt_name = "Planilha_Mestra.xlsx" if self.caminho_mestra.name == "planilha_mestra.xlsx" else "planilha_mestra.xlsx"
                alt_path = self.caminho_mestra.parent / alt_name
                if alt_path.exists():
                    self.caminho_mestra = alt_path
                else:
                    raise FileNotFoundError(
                        f"Planilha mestra não encontrada: "
                        f"{self.caminho_mestra}"
                    )

            self.wb = load_workbook(
                self.caminho_mestra
            )

            self.ws = self.wb.active
        except Exception as erro:
            raise ErroPlanilhaMestra(
                f"Não foi possível acessar a planilha mestra: {erro}"
            ) from erro
        
    def localizar_colunas(self):

        self.colunas = {}

        for coluna in range(
            1,
            self.ws.max_column + 1
        ):

            nome = self.ws.cell(
                1,
                coluna
            ).value

            if nome:
                self.colunas[
                    str(nome).strip()
                ] = coluna

    def ler_clientes(self):

        clientes = []

        for linha in range(
            2,
            self.ws.max_row + 1
        ):

            cliente = {}

            for nome, coluna in self.colunas.items():

                cliente[nome] = self.ws.cell(
                    linha,
                    coluna
                ).value

            # Verifica se a linha está completamente vazia
            possui_dados = any(
                valor is not None
                and str(valor).strip() != ""
                for valor in cliente.values()
            )

            if not possui_dados:
                continue

            # Linha possui dados, mas não possui protocolo
            protocolo = cliente.get("Protocolo")

            if protocolo is None or str(protocolo).strip() == "":
                logger.warning(
                    "Linha %s possui dados, mas não possui protocolo. Ignorando.",
                    linha
                )
                continue

            clientes.append(cliente)

        return clientes

    def definir_atendimento(self, cliente):

        status = str(cliente.get("Status") or "").strip()

        if status in ("ATIVO", "CONCLUIDO_P2", "CONCLUIDO_P3", "CADASTRADO", "CONCLUIDO"):

            return {
                "tipo": "CADASTRO_OK",
                "mensagem": (
                    "Seu cadastro foi concluído "
                    "com sucesso."
                )
            }

        return None

    def comunicar(self, cliente, atendimento):

        print()
        print("=" * 60)
        print("COMUNICAÇÃO SAC")
        print("=" * 60)

        print(f"Cliente   : {cliente.get('Nome')}")
        print(f"CPF       : {cliente.get('CPF')}")
        print(f"E-mail    : {cliente.get('E-mail')}")
        print(f"Protocolo : {cliente.get('Protocolo')}")
        print(f"Tipo      : {atendimento['tipo']}")
        print()
        print(atendimento["mensagem"])

        print("=" * 60)

        logger.info(
            "Comunicação preparada para protocolo %s",
            cliente.get("Protocolo")
        )

    def processar_cliente(self, cliente):

        protocolo = cliente.get("Protocolo")

        if not protocolo:

            logger.warning(
                "Registro sem protocolo. Ignorando."
            )

            return

        logger.info(
            "Processando protocolo: %s",
            protocolo
        )

        # ----------------------------------------------------
        # 1. Verificar se o status gera atendimento
        # ----------------------------------------------------

        atendimento = self.definir_atendimento(
            cliente
        )

        if atendimento is None:

            logger.info(
                "Status '%s' não possui atendimento "
                "configurado.",
                cliente.get("Status")
            )

            return

        # ----------------------------------------------------
        # 2. Criar chave do atendimento
        # ----------------------------------------------------

        chave = (
            f"{protocolo}_{atendimento['tipo']}"
        )

        # ----------------------------------------------------
        # 3. Verificar duplicidade
        # ----------------------------------------------------

        if self.registro.ja_processado(chave):

            logger.info(
                "Atendimento já realizado: %s",
                chave
            )

            return

        # ----------------------------------------------------
        # 4. Comunicar cliente
        # ----------------------------------------------------
        data_atendimento = None
        
        try:

            data_atendimento = enviar_comunicacao(
                cliente
            )

            self.registro.registrar(
                cliente=cliente,
                tipo_atendimento=atendimento["tipo"],
                status_atendimento="COMUNICADO",
                data_atendimento=data_atendimento,
                observacao=(
                    "E-mail enviado com sucesso."
                )
            )

        except Exception as erro:

            logger.error(
                "Erro ao comunicar cliente %s: %s",
                cliente.get("CPF"),
                erro
            )

            self.registro.registrar(
                cliente=cliente,
                tipo_atendimento=atendimento["tipo"],
                status_atendimento="PENDENTE_CONTATO",
                data_atendimento=None,
                observacao=str(erro)
            )
        return

    def executar(self):

        logger.info(
            "Iniciando Processo 4 - SAC"
        )

        self.registro.carregar()

        try:
            self.carregar_planilha_mestra()
        except Exception as erro:
            logger.error(
                "Falha no acesso à planilha mestra: %s",
                erro
            )
            return {
                "sucesso": False,
                "erro": str(erro),
                "total_processados": 0,
                "total_registros": 0
            }
        
        self.localizar_colunas()

        clientes = self.ler_clientes()

        logger.info(
            "Total de registros encontrados: %d",
            len(clientes)
        )

        processados = 0
        for cliente in clientes:

            self.processar_cliente(
                cliente
            )
            processados += 1

        logger.info(
            "Processo 4 - SAC finalizado"
        )
        return {
            "sucesso": True,
            "total_processados": processados,
            "total_registros": len(clientes)
        }


def executar_processo_sac(caminho_planilha_mestra=None, caminho_planilha_sac=None):
    """Função utilitária para execução programática do Processo 4 (SAC)."""
    processo = ProcessoSAC(
        caminho_planilha_mestra=caminho_planilha_mestra,
        caminho_planilha_sac=caminho_planilha_sac
    )
    return processo.executar()


# ============================================================
# EXECUÇÃO
# ============================================================

if __name__ == "__main__":

    processo = ProcessoSAC(
        caminho_planilha_mestra=PLANILHA_MESTRA,
        caminho_planilha_sac=PLANILHA_SAC
    )

    processo.executar()