"""
Módulo de Consolidação de Dados - Processo 5 (Relatórios e Gerência)
Responsável por carregar e cruzar os dados da Planilha Mestra e da Planilha SAC.
"""
from pathlib import Path
from typing import Dict, List, Any, Optional
from openpyxl import load_workbook


def get_project_root() -> Path:
    current = Path(__file__).resolve().parent
    for parent in [current] + list(current.parents):
        if (parent / "ERP_Portal_Fake").exists() or (parent / "HyperAutomation").exists():
            return parent
    return Path(__file__).resolve().parents[3]


class ConsolidadorDados:
    """Consolida os dados dos Processos 1 a 4 a partir das planilhas do sistema."""

    def __init__(self, caminho_mestra: Optional[Path] = None, caminho_sac: Optional[Path] = None):
        root = get_project_root()
        dir_sistema = root / "ERP_Portal_Fake" / "Sistema_Integrador_Portal_Fake"

        # Resolução automática de case sensitivity para Planilha Mestra
        if caminho_mestra:
            self.caminho_mestra = Path(caminho_mestra)
        else:
            p1 = dir_sistema / "Planilha_Mestra.xlsx"
            p2 = dir_sistema / "planilha_mestra.xlsx"
            self.caminho_mestra = p1 if p1.exists() else p2

        if caminho_sac:
            self.caminho_sac = Path(caminho_sac)
        else:
            self.caminho_sac = dir_sistema / "atendimentos_sac.xlsx"

    def carregar_planilha_mestra(self, caminho: Optional[Path] = None) -> List[Dict[str, Any]]:
        """Lê os registros da Planilha Mestra."""
        target_path = Path(caminho) if caminho else self.caminho_mestra

        if not target_path.exists():
            alt_name = "Planilha_Mestra.xlsx" if target_path.name == "planilha_mestra.xlsx" else "planilha_mestra.xlsx"
            alt_path = target_path.parent / alt_name
            if alt_path.exists():
                target_path = alt_path
            else:
                return []

        try:
            wb = load_workbook(target_path, data_only=True)
            ws = wb.active

            if ws.max_row < 2:
                return []

            # Mapeia colunas do cabeçalho
            colunas = {}
            for col in range(1, ws.max_column + 1):
                val = ws.cell(1, col).value
                if val:
                    colunas[str(val).strip()] = col

            registros = []
            for row in range(2, ws.max_row + 1):
                item = {}
                for nome, col_idx in colunas.items():
                    item[nome] = ws.cell(row, col_idx).value

                # Verifica se há dados na linha
                tem_dados = any(v is not None and str(v).strip() != "" for v in item.values())
                if not tem_dados:
                    continue

                cpf = str(item.get("CPF") or "").strip()
                cpf_digs = "".join(filter(str.isdigit, cpf))
                if len(cpf_digs) < 11 and len(cpf_digs) > 0:
                    cpf_digs = cpf_digs.zfill(11)

                item["cpf_limpo"] = cpf_digs
                item["linha_planilha"] = row
                registros.append(item)

            return registros
        except Exception:
            return []

    def carregar_planilha_sac(self, caminho: Optional[Path] = None) -> List[Dict[str, Any]]:
        """Lê os registros da planilha de Atendimentos SAC."""
        target_path = Path(caminho) if caminho else self.caminho_sac

        if not target_path or not target_path.exists():
            return []

        try:
            wb = load_workbook(target_path, data_only=True)
            ws = wb.active

            if ws.max_row < 2:
                return []

            colunas = {}
            for col in range(1, ws.max_column + 1):
                val = ws.cell(1, col).value
                if val:
                    colunas[str(val).strip()] = col

            registros = []
            for row in range(2, ws.max_row + 1):
                item = {}
                for nome, col_idx in colunas.items():
                    item[nome] = ws.cell(row, col_idx).value

                tem_dados = any(v is not None and str(v).strip() != "" for v in item.values())
                if not tem_dados:
                    continue

                cpf = str(item.get("CPF") or "").strip()
                cpf_digs = "".join(filter(str.isdigit, cpf))
                if len(cpf_digs) < 11 and len(cpf_digs) > 0:
                    cpf_digs = cpf_digs.zfill(11)

                item["cpf_limpo"] = cpf_digs
                registros.append(item)

            return registros
        except Exception:
            return []

    def consolidar(self) -> Dict[str, Any]:
        """
        Consolida a visão unificada dos clientes cruzando a Planilha Mestra e o SAC.
        """
        mestra_dados = self.carregar_planilha_mestra()
        sac_dados = self.carregar_planilha_sac()

        # Indexa SAC por CPF e Protocolo
        sac_por_cpf = {}
        sac_por_protocolo = {}
        for s in sac_dados:
            cpf_limpo = s.get("cpf_limpo")
            prot = str(s.get("Protocolo") or "").strip()
            if cpf_limpo:
                sac_por_cpf[cpf_limpo] = s
            if prot:
                sac_por_protocolo[prot] = s

        clientes_consolidados = []
        cpfs_processados = set()

        for m in mestra_dados:
            cpf_limpo = m.get("cpf_limpo", "")
            prot = str(m.get("Protocolo") or "").strip()

            sac_info = sac_por_cpf.get(cpf_limpo) or sac_por_protocolo.get(prot) or {}

            status_cadastro = str(m.get("Status") or "PENDENTE").strip()
            status_sac = str(sac_info.get("Status Atendimento") or "NÃO INICIADO").strip()

            # Classificação da etapa concluída
            if status_sac in ["COMUNICADO"]:
                fase_concluida = "PROCESSO 4 (SAC CONCLUÍDO)"
            elif status_cadastro in ["CONCLUIDO_P3", "CADASTRADO", "CONCLUIDO", "ATIVO"]:
                fase_concluida = "PROCESSO 3 (CADASTRO CONCLUÍDO)"
            elif status_cadastro in ["CONCLUIDO_P2"]:
                fase_concluida = "PROCESSO 2 (DADOS ORGANIZADOS)"
            elif status_cadastro in ["DUPLICADO", "DUPLICADO_P3"]:
                fase_concluida = "PROCESSO 3 (DUPLICADO)"
            elif "ERRO" in status_cadastro or "REJEITADO" in status_cadastro:
                fase_concluida = f"PROCESSO 3 ({status_cadastro})"
            else:
                fase_concluida = "PROCESSO 1 (EM ANDAMENTO)"

            consolidado = {
                "cpf": m.get("CPF") or (sac_info.get("CPF") if sac_info else ""),
                "cpf_limpo": cpf_limpo,
                "protocolo": prot or (sac_info.get("Protocolo") if sac_info else "N/A"),
                "nome": m.get("Nome") or (sac_info.get("Nome") if sac_info else "N/A"),
                "email": m.get("E-mail") or (sac_info.get("E-mail") if sac_info else ""),
                "telefone": m.get("Telefone") or "",
                "endereco": m.get("Endereço") or "",
                "nascimento": m.get("Data de Nascimento") or "",
                "status_cadastro": status_cadastro,
                "data_processamento_cadastro": m.get("Data de Processamento") or "",
                "observacoes_cadastro": m.get("Observações") or "",
                "status_sac": status_sac,
                "tipo_atendimento_sac": sac_info.get("Tipo Atendimento") or "N/A",
                "data_atendimento_sac": sac_info.get("Data Atendimento") or "",
                "observacao_sac": sac_info.get("Observação") or "",
                "fase_concluida": fase_concluida
            }
            clientes_consolidados.append(consolidado)
            if cpf_limpo:
                cpfs_processados.add(cpf_limpo)

        # Inclui registros de SAC que porventura não estejam na Mestra
        for s in sac_dados:
            cpf_limpo = s.get("cpf_limpo", "")
            if cpf_limpo and cpf_limpo not in cpfs_processados:
                clientes_consolidados.append({
                    "cpf": s.get("CPF") or "",
                    "cpf_limpo": cpf_limpo,
                    "protocolo": s.get("Protocolo") or "N/A",
                    "nome": s.get("Nome") or "N/A",
                    "email": s.get("E-mail") or "",
                    "telefone": "",
                    "endereco": "",
                    "nascimento": "",
                    "status_cadastro": s.get("Status Cadastro") or "CONCLUIDO",
                    "data_processamento_cadastro": "",
                    "observacoes_cadastro": "",
                    "status_sac": s.get("Status Atendimento") or "NÃO INICIADO",
                    "tipo_atendimento_sac": s.get("Tipo Atendimento") or "N/A",
                    "data_atendimento_sac": s.get("Data Atendimento") or "",
                    "observacao_sac": s.get("Observação") or "",
                    "fase_concluida": "PROCESSO 4 (SAC CONCLUÍDO)" if s.get("Status Atendimento") == "COMUNICADO" else "PROCESSO 4 (SAC PENDENTE)"
                })
                cpfs_processados.add(cpf_limpo)

        return {
            "total_clientes": len(clientes_consolidados),
            "clientes": clientes_consolidados,
            "total_mestra": len(mestra_dados),
            "total_sac": len(sac_dados)
        }
