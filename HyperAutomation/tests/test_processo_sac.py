from pathlib import Path
from unittest.mock import MagicMock, patch

import pytest
from openpyxl import Workbook, load_workbook

from source.processo_atendimento.processo_sac import (
    ErroPlanilhaMestra,
    ProcessoSAC,
    RegistroAtendimento,
    criar_mensagem,
    enviar_comunicacao,
)

# ============================================================
# DADOS DE TESTE
# ============================================================

@pytest.fixture
def cliente():
    return {
        "Protocolo": "PROT-20260820-001",
        "CPF": "10000013441",
        "Nome": "Hugo Lima",
        "E-mail": "hugo@example.com",
        "Status": "ATIVO",
    }


@pytest.fixture
def cliente_sem_email():
    return {
        "Protocolo": "PROT-20260820-002",
        "CPF": "10000013442",
        "Nome": "Maria Silva",
        "E-mail": None,
        "Status": "ATIVO",
    }


@pytest.fixture
def processo(tmp_path):
    return ProcessoSAC(
        caminho_planilha_mestra=tmp_path / "planilha_mestra.xlsx",
        caminho_planilha_sac=tmp_path / "atendimentos_sac.xlsx",
    )


# ============================================================
# REGISTRO DE ATENDIMENTO
# ============================================================

def test_criar_planilha_atendimento(tmp_path):

    caminho = tmp_path / "atendimentos_sac.xlsx"

    registro = RegistroAtendimento(caminho)

    registro.criar_planilha()

    assert caminho.exists()

    wb = load_workbook(caminho)
    ws = wb["Atendimentos"]

    cabecalhos = [
        ws.cell(1, coluna).value
        for coluna in range(1, ws.max_column + 1)
    ]

    assert cabecalhos == [
        "Chave Atendimento",
        "Protocolo",
        "CPF",
        "Nome",
        "E-mail",
        "Status Cadastro",
        "Tipo Atendimento",
        "Status Atendimento",
        "Data Atendimento",
        "Observação",
    ]


def test_criar_planilha_nao_recria_arquivo_existente(tmp_path):

    caminho = tmp_path / "atendimentos_sac.xlsx"

    registro = RegistroAtendimento(caminho)

    registro.criar_planilha()

    primeira_data = caminho.stat().st_mtime

    registro.criar_planilha()

    segunda_data = caminho.stat().st_mtime

    assert caminho.exists()
    assert segunda_data == primeira_data


def test_registrar_atendimento(tmp_path, cliente):

    caminho = tmp_path / "atendimentos_sac.xlsx"

    registro = RegistroAtendimento(caminho)
    registro.carregar()

    registro.registrar(
        cliente=cliente,
        tipo_atendimento="CADASTRO_OK",
        status_atendimento="COMUNICADO",
        data_atendimento="20/08/2026 10:30",
        observacao="E-mail enviado com sucesso.",
    )

    wb = load_workbook(caminho)
    ws = wb["Atendimentos"]

    assert ws.max_row == 2

    assert ws.cell(2, 1).value == (
        "PROT-20260820-001_CADASTRO_OK"
    )

    assert ws.cell(2, 2).value == "PROT-20260820-001"
    assert ws.cell(2, 3).value == "10000013441"
    assert ws.cell(2, 4).value == "Hugo Lima"
    assert ws.cell(2, 8).value == "COMUNICADO"


def test_ja_processado_retorna_true_para_atendimento_existente(
    tmp_path,
    cliente,
):

    caminho = tmp_path / "atendimentos_sac.xlsx"

    registro = RegistroAtendimento(caminho)
    registro.carregar()

    registro.registrar(
        cliente=cliente,
        tipo_atendimento="CADASTRO_OK",
        status_atendimento="COMUNICADO",
        data_atendimento="20/08/2026 10:30",
        observacao="Teste",
    )

    resultado = registro.ja_processado(
        "PROT-20260820-001_CADASTRO_OK"
    )

    assert resultado is True


def test_ja_processado_retorna_false_para_atendimento_inexistente(
    tmp_path,
):

    caminho = tmp_path / "atendimentos_sac.xlsx"

    registro = RegistroAtendimento(caminho)
    registro.carregar()

    resultado = registro.ja_processado(
        "PROT-NAO-EXISTE_CADASTRO_OK"
    )

    assert resultado is False


# ============================================================
# CRIAÇÃO DA MENSAGEM
# ============================================================

def test_criar_mensagem(cliente):

    assunto, mensagem, data = criar_mensagem(cliente)

    assert assunto == (
        "Cadastro realizado com sucesso - "
        "Protocolo PROT-20260820-001"
    )

    assert "Olá, Hugo Lima!" in mensagem
    assert "seu cadastro foi realizado com sucesso" in mensagem
    assert "PROT-20260820-001" in mensagem
    assert "Data do atendimento:" in mensagem

    assert data is not None


def test_criar_mensagem_retorna_data_no_formato_esperado(cliente):

    _, _, data = criar_mensagem(cliente)

    partes = data.split(" ")

    assert len(partes) == 2

    data_parte = partes[0]
    hora_parte = partes[1]

    assert len(data_parte.split("/")) == 3
    assert len(hora_parte.split(":")) == 2


# ============================================================
# ENVIO DE COMUNICAÇÃO
# ============================================================

def test_enviar_comunicacao_sem_email(cliente_sem_email):

    with pytest.raises(
        ValueError,
        match="não possui e-mail",
    ):
        enviar_comunicacao(cliente_sem_email)


def test_enviar_comunicacao_com_sucesso(cliente):

    servidor_mock = MagicMock()
    servidor_mock.__enter__.return_value = servidor_mock

    with patch(
        "source.processo_atendimento.processo_sac.smtplib.SMTP",
        return_value=servidor_mock,
    ):

        data = enviar_comunicacao(cliente)

    servidor_mock.starttls.assert_called_once()

    servidor_mock.login.assert_called_once()

    servidor_mock.send_message.assert_called_once()

    assert data is not None


def test_enviar_comunicacao_falha_no_smtp(cliente):

    servidor_mock = MagicMock()
    servidor_mock.__enter__.return_value = servidor_mock
    servidor_mock.starttls.side_effect = ConnectionError("Erro SMTP")

    with patch(
        "source.processo_atendimento.processo_sac.smtplib.SMTP",
        return_value=servidor_mock,
    ):

        with pytest.raises(ConnectionError):
            with patch(
                "source.processo_atendimento.processo_sac.smtplib.SMTP",
                return_value=servidor_mock,
            ):
                enviar_comunicacao(cliente)

# ============================================================
# PLANILHA MESTRA
# ============================================================

def criar_planilha_mestra(caminho, linhas):

    wb = Workbook()
    ws = wb.active

    ws.append([
        "Protocolo",
        "CPF",
        "Nome",
        "E-mail",
        "Status",
    ])

    for linha in linhas:
        ws.append(linha)

    wb.save(caminho)


def test_carregar_planilha_mestra_com_sucesso(
    processo,
    tmp_path,
):

    caminho = tmp_path / "planilha_mestra.xlsx"

    criar_planilha_mestra(
        caminho,
        [
            [
                "PROT-001",
                "11111111111",
                "João Silva",
                "joao@example.com",
                "ATIVO",
            ]
        ],
    )

    processo.caminho_mestra = caminho

    processo.carregar_planilha_mestra()

    assert processo.wb is not None
    assert processo.ws is not None


def test_carregar_planilha_mestra_inexistente(processo):

    with pytest.raises(ErroPlanilhaMestra):

        processo.carregar_planilha_mestra()


def test_carregar_planilha_mestra_erro_do_openpyxl(
    processo,
    tmp_path,
):

    caminho = tmp_path / "planilha_mestra.xlsx"

    caminho.write_text(
        "arquivo inválido",
        encoding="utf-8",
    )

    processo.caminho_mestra = caminho

    with pytest.raises(ErroPlanilhaMestra):

        processo.carregar_planilha_mestra()


# ============================================================
# LOCALIZAÇÃO DAS COLUNAS
# ============================================================

def test_localizar_colunas(processo, tmp_path):

    caminho = tmp_path / "planilha_mestra.xlsx"

    criar_planilha_mestra(
        caminho,
        [],
    )

    processo.caminho_mestra = caminho

    processo.carregar_planilha_mestra()
    processo.localizar_colunas()

    assert processo.colunas["Protocolo"] == 1
    assert processo.colunas["CPF"] == 2
    assert processo.colunas["Nome"] == 3
    assert processo.colunas["E-mail"] == 4
    assert processo.colunas["Status"] == 5


def test_localizar_colunas_remove_espacos_do_cabecalho(
    processo,
    tmp_path,
):

    caminho = tmp_path / "planilha_mestra.xlsx"

    wb = Workbook()
    ws = wb.active

    ws.append([
        " Protocolo ",
        " CPF",
        "Nome ",
    ])

    wb.save(caminho)

    processo.caminho_mestra = caminho

    processo.carregar_planilha_mestra()
    processo.localizar_colunas()

    assert "Protocolo" in processo.colunas
    assert "CPF" in processo.colunas
    assert "Nome" in processo.colunas


# ============================================================
# LEITURA DOS CLIENTES
# ============================================================

def test_ler_clientes(processo, cliente, tmp_path):

    caminho = tmp_path / "planilha_mestra.xlsx"

    criar_planilha_mestra(
        caminho,
        [
            [
                cliente["Protocolo"],
                cliente["CPF"],
                cliente["Nome"],
                cliente["E-mail"],
                cliente["Status"],
            ]
        ],
    )

    processo.caminho_mestra = caminho

    processo.carregar_planilha_mestra()
    processo.localizar_colunas()

    clientes = processo.ler_clientes()

    assert len(clientes) == 1

    assert clientes[0]["Protocolo"] == (
        "PROT-20260820-001"
    )

    assert clientes[0]["Nome"] == "Hugo Lima"


def test_ler_clientes_ignora_linha_completamente_vazia(
    processo,
    tmp_path,
):

    caminho = tmp_path / "planilha_mestra.xlsx"

    criar_planilha_mestra(
        caminho,
        [
            [
                "PROT-001",
                "11111111111",
                "João",
                "joao@example.com",
                "ATIVO",
            ],
            [
                None,
                None,
                None,
                None,
                None,
            ],
        ],
    )

    processo.caminho_mestra = caminho

    processo.carregar_planilha_mestra()
    processo.localizar_colunas()

    clientes = processo.ler_clientes()

    assert len(clientes) == 1


def test_ler_clientes_ignora_registro_sem_protocolo(
    processo,
    tmp_path,
):

    caminho = tmp_path / "planilha_mestra.xlsx"

    criar_planilha_mestra(
        caminho,
        [
            [
                None,
                "11111111111",
                "João",
                "joao@example.com",
                "ATIVO",
            ],
        ],
    )

    processo.caminho_mestra = caminho

    processo.carregar_planilha_mestra()
    processo.localizar_colunas()

    clientes = processo.ler_clientes()

    assert clientes == []


def test_ler_clientes_ignora_protocolo_apenas_com_espacos(
    processo,
    tmp_path,
):

    caminho = tmp_path / "planilha_mestra.xlsx"

    criar_planilha_mestra(
        caminho,
        [
            [
                "   ",
                "11111111111",
                "João",
                "joao@example.com",
                "ATIVO",
            ],
        ],
    )

    processo.caminho_mestra = caminho

    processo.carregar_planilha_mestra()
    processo.localizar_colunas()

    clientes = processo.ler_clientes()

    assert clientes == []


# ============================================================
# DEFINIÇÃO DO ATENDIMENTO
# ============================================================

@pytest.mark.parametrize(
    "status",
    [
        "ATIVO",
        "CONCLUIDO_P2",
    ],
)
def test_definir_atendimento_status_que_gera_cadastro_ok(
    processo,
    status,
):

    cliente = {
        "Status": status
    }

    atendimento = processo.definir_atendimento(cliente)

    assert atendimento is not None
    assert atendimento["tipo"] == "CADASTRO_OK"

    assert atendimento["mensagem"] == (
        "Seu cadastro foi concluído com sucesso."
    )


@pytest.mark.parametrize(
    "status",
    [
        None,
        "",
        "INATIVO",
        "PENDENTE",
        "CANCELADO",
        "QUALQUER_OUTRO_STATUS",
    ],
)
def test_definir_atendimento_status_sem_configuracao(
    processo,
    status,
):

    cliente = {
        "Status": status
    }

    atendimento = processo.definir_atendimento(cliente)

    assert atendimento is None


# ============================================================
# COMUNICAR
# ============================================================

def test_comunicar_nao_envia_email(
    processo,
    cliente,
    capsys,
):

    atendimento = processo.definir_atendimento(cliente)

    processo.comunicar(
        cliente,
        atendimento,
    )

    saida = capsys.readouterr().out

    assert "COMUNICAÇÃO SAC" in saida
    assert "Hugo Lima" in saida
    assert "10000013441" in saida
    assert "PROT-20260820-001" in saida
    assert "CADASTRO_OK" in saida
    assert (
        "Seu cadastro foi concluído com sucesso."
        in saida
    )


# ============================================================
# PROCESSAMENTO DO CLIENTE
# ============================================================

def test_processar_cliente_sem_protocolo(
    processo,
):

    cliente = {
        "Protocolo": None,
        "CPF": "11111111111",
        "Nome": "João",
        "E-mail": "joao@example.com",
        "Status": "ATIVO",
    }

    processo.registro = MagicMock()

    processo.processar_cliente(cliente)

    processo.registro.ja_processado.assert_not_called()
    processo.registro.registrar.assert_not_called()


def test_processar_cliente_status_sem_atendimento(
    processo,
):

    cliente = {
        "Protocolo": "PROT-001",
        "CPF": "11111111111",
        "Nome": "João",
        "E-mail": "joao@example.com",
        "Status": "INATIVO",
    }

    processo.registro = MagicMock()

    processo.processar_cliente(cliente)

    processo.registro.ja_processado.assert_not_called()
    processo.registro.registrar.assert_not_called()


def test_processar_cliente_ja_processado(
    processo,
    cliente,
):

    processo.registro = MagicMock()

    processo.registro.ja_processado.return_value = True

    with patch(
        "source.processo_atendimento.processo_sac.enviar_comunicacao"
    ) as enviar_mock:

        processo.processar_cliente(cliente)

    enviar_mock.assert_not_called()
    processo.registro.registrar.assert_not_called()


def test_processar_cliente_comunicacao_sucesso(
    processo,
    cliente,
):

    processo.registro = MagicMock()

    processo.registro.ja_processado.return_value = False


    with patch(
        "source.processo_atendimento.processo_sac.enviar_comunicacao",
        return_value="20/08/2026 10:30"
    ) as enviar_mock:

        processo.processar_cliente(cliente)

    enviar_mock.assert_called_once_with(cliente)

    processo.registro.registrar.assert_called_once()

    argumentos = (
        processo.registro.registrar.call_args.kwargs
    )

    assert argumentos["cliente"] == cliente
    assert argumentos["tipo_atendimento"] == "CADASTRO_OK"
    assert argumentos["status_atendimento"] == "COMUNICADO"
    assert argumentos["data_atendimento"] == (
        "20/08/2026 10:30"
    )
    assert argumentos["observacao"] == (
        "E-mail enviado com sucesso."
    )


def test_processar_cliente_falha_comunicacao(
    processo,
    cliente_sem_email,
):

    processo.registro = MagicMock()

    processo.registro.ja_processado.return_value = False

    with patch(
        "source.processo_atendimento.processo_sac.enviar_comunicacao",
        side_effect=ConnectionError(
            "Cliente Maria Silva não possui e-mail."
        ),
    ) as enviar_mock:

        processo.processar_cliente(
            cliente_sem_email
        )

    enviar_mock.assert_called_once_with(
        cliente_sem_email
    )

    processo.registro.registrar.assert_called_once()

    argumentos = (
        processo.registro.registrar.call_args.kwargs
    )

    assert argumentos["cliente"] == cliente_sem_email
    assert argumentos["tipo_atendimento"] == "CADASTRO_OK"
    assert argumentos["status_atendimento"] == (
        "PENDENTE_CONTATO"
    )
    assert argumentos["data_atendimento"] is None
    assert "não possui e-mail" in (
        argumentos["observacao"]
    )


# ============================================================
# EXECUÇÃO DO PROCESSO
# ============================================================

def test_executar_para_quando_planilha_mestra_nao_existe(
    processo,
):

    processo.registro = MagicMock()

    with patch.object(
        processo,
        "localizar_colunas",
    ) as localizar_mock:

        processo.executar()

    localizar_mock.assert_not_called()


def test_executar_processa_clientes(
    processo,
    cliente,
):

    processo.registro = MagicMock()

    with patch.object(
        processo,
        "carregar_planilha_mestra",
    ) as carregar_mock, patch.object(
        processo,
        "localizar_colunas",
    ) as localizar_mock, patch.object(
        processo,
        "ler_clientes",
        return_value=[cliente],
    ) as ler_mock, patch.object(
        processo,
        "processar_cliente",
    ) as processar_mock:

        processo.executar()

    processo.registro.carregar.assert_called_once()

    carregar_mock.assert_called_once()
    localizar_mock.assert_called_once()
    ler_mock.assert_called_once()

    processar_mock.assert_called_once_with(
        cliente
    )


def test_executar_processa_varios_clientes(
    processo,
):

    clientes = [
        {
            "Protocolo": "PROT-001",
            "CPF": "11111111111",
            "Nome": "João",
            "E-mail": "joao@example.com",
            "Status": "ATIVO",
        },
        {
            "Protocolo": "PROT-002",
            "CPF": "22222222222",
            "Nome": "Maria",
            "E-mail": "maria@example.com",
            "Status": "CONCLUIDO_P2",
        },
    ]

    processo.registro = MagicMock()

    with patch.object(
        processo,
        "carregar_planilha_mestra",
    ), patch.object(
        processo,
        "localizar_colunas",
    ), patch.object(
        processo,
        "ler_clientes",
        return_value=clientes,
    ), patch.object(
        processo,
        "processar_cliente",
    ) as processar_mock:

        processo.executar()

    assert processar_mock.call_count == 2

    processar_mock.assert_any_call(
        clientes[0]
    )

    processar_mock.assert_any_call(
        clientes[1]
    )