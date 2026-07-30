import os
import sys
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def build_docx():
    doc = docx.Document()

    # Define Margins (2.5 cm)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.98)
        section.bottom_margin = Inches(0.98)
        section.left_margin = Inches(0.98)
        section.right_margin = Inches(0.98)

    # Styles
    style_normal = doc.styles['Normal']
    font_normal = style_normal.font
    font_normal.name = 'Calibri'
    font_normal.size = Pt(11)
    font_normal.color.rgb = RGBColor(0x22, 0x22, 0x22)

    # Colors
    NAVY = RGBColor(0x1B, 0x36, 0x5D)
    DARK_BLUE = RGBColor(0x0F, 0x4C, 0x81)
    GRAY = RGBColor(0x55, 0x55, 0x55)

    def add_title(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(20)
        run.font.color.rgb = NAVY
        p.paragraph_format.space_after = Pt(4)
        return p

    def add_subtitle(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.font.size = Pt(13)
        run.font.color.rgb = GRAY
        run.italic = True
        p.paragraph_format.space_after = Pt(18)
        return p

    def add_h1(text):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(15)
        run.font.color.rgb = NAVY
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(12)
        run.font.color.rgb = DARK_BLUE
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        return p

    def add_p(text, bold_prefix=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.bold = True
            r_bold.font.color.rgb = NAVY
        p.add_run(text)
        return p

    def add_bullet(text, bold_prefix=None):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.bold = True
            r_bold.font.color.rgb = NAVY
        p.add_run(text)
        return p

    def add_code_block(code_text):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F4F6F8"/>')
        tcPr.append(shd)
        
        # Borders
        borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="single" w:sz="4" w:space="0" w:color="D0D7DE"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="D0D7DE"/><w:left w:val="single" w:sz="12" w:space="0" w:color="0F4C81"/><w:right w:val="single" w:sz="4" w:space="0" w:color="D0D7DE"/></w:tcBorders>')
        tcPr.append(borders)

        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(code_text)
        run.font.name = 'Courier New'
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # HEADER & IDENTIFICATION
    add_title("Relatório Técnico – Automação do Processo 1")
    add_subtitle("Empresa Portal Fake | Roteiro 10 – Técnicas de Hyperautomation (Semana 07)")

    add_p("Prof. Moisés Levy", bold_prefix="Professor: ")
    add_p("Técnicas de Hyperautomation / PÓLO DE INOVAÇÃO IFAM - FAEPI", bold_prefix="Disciplina/Curso: ")
    add_p("Eric Luna Costa, Daniele Greice Albuquerque e Silva, Kauã Sales Viana, Sannyer Cardoso Carvalho Nery", bold_prefix="Alunos/Integrantes: ")

    # SECTION 1: IDENTIFICAÇÃO DO PROJETO E EQUIPE
    add_h1("1. Identificação do Projeto e Organização da Equipe")
    add_p("O objetivo deste módulo foi desenvolver a automação fim a fim do Processo 1 (Setor de Atendimento) da Empresa Portal Fake. Em conformidade com as diretrizes do Roteiro 10, o robô foi estruturado não como uma aplicação isolada, mas como um módulo perfeitamente integrado à solução de Hyperautomation preexistente, preservando a arquitetura limpa, o histórico de versionamento no Git e o padrão de projeto.")

    add_h2("1.1 Divisão de Responsabilidades da Equipe")
    add_p("Para garantir cobertura integral dos requisitos do projeto e eficiência na execução, as atribuições foram distribuídas de forma multifuncional entre os membros da equipe:")

    # Table of Team Members
    table = doc.add_table(rows=5, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Integrante"
    hdr_cells[1].text = "Responsabilidade e Papel no Projeto"
    for cell in hdr_cells:
        tcPr = cell._tc.get_or_add_tcPr()
        tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="1B365D"/>'))
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    team_data = [
        ("Eric Luna Costa", "Analista de Processos BPMN & Desenvolvimento Core (Modelagem BPMN no Draw.io, validação documental e regras de negócio)."),
        ("Daniele Greice Albuquerque e Silva", "Desenvolvimento de Automação (Auxílio na lógica de captura de dados e módulos de suporte)."),
        ("Kauã Sales Viana", "Desenvolvimento de Automação (Testes de módulos e apoio no fluxo de captura web)."),
        ("Sannyer Cardoso Carvalho Nery", "DevOps, Versionamento (GitFlow), Arquitetura da Automação (Integração Google Drive API v3, Gestor de Arquivos ERP, Notificações HTML, Orquestrador Geral e Redação do Relatório Técnico).")
    ]

    for idx, (name, resp) in enumerate(team_data, start=1):
        row_cells = table.rows[idx].cells
        row_cells[0].text = name
        row_cells[1].text = resp
        row_cells[0].paragraphs[0].runs[0].bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # SECTION 2: DESCRIÇÃO DA SOLUÇÃO DESENVOLVIDA
    add_h1("2. Descrição da Solução Desenvolvida")
    add_p("A automação do Setor de Atendimento resolve a triagem e o processamento de solicitações de clientes através de um ciclo de 6 etapas autônomas:")

    add_bullet(" Monitora a caixa de entrada (IMAP4_SSL) por novas solicitações não lidas de clientes contendo 'Assinatura' ou 'Ficha' no assunto.", bold_prefix="1. Recebimento: ")
    add_bullet(" Baixa automaticamente o arquivo anexado (PDF Único contendo Ficha Assinada + Documentos com Foto + Comprovante de Residência) para a pasta de entrada 'Downloads'.", bold_prefix="2. Download: ")
    add_bullet(" Executa inspeção de conteúdo no arquivo PDF via módulo 'pypdf', verificando a integridade do documento e confirmando se a Ficha Cadastral está assinada e acompanhada dos documentos obrigatórios.", bold_prefix="3. Validação: ")
    add_bullet(" Classifica o arquivo e o move para 'Documentos_OK' (se aprovado) ou 'Documentos_Pendentes' (se reprovado/incompleto), sincronizando instantaneamente a movimentação na pasta física local e na nuvem via Google Drive API v3.", bold_prefix="4. Classificação: ")
    add_bullet(" Dispara e-mail transacional automatizado em HTML responsivo ao cliente informando o status (Confirmação de Aprovação ou Notificação de Pendência com orientações de reenvio).", bold_prefix="5. Retorno: ")
    add_bullet(" Interage com o Portal Fake ERP via Playwright para efetuar o cadastro definitivo do cliente e transfere o PDF aprovado para a pasta 'Encaminhados' (próximo setor).", bold_prefix="6. Encaminhamento: ")

    add_h2("2.1 Integração com o Google Drive API v3 e Resolução de Cota")
    add_p("Uma das contribuições arquiteturais de maior destaque nesta versão V2.0 foi o desenvolvimento do módulo ", bold_prefix="Destaque Técnico: ")
    p_drive = doc.paragraphs[-1]
    p_drive.add_run("GestorDrive (gestor_drive.py). O módulo conecta à API do Google Drive v3, garante a criação e mapeamento das subpastas remotas (Downloads, Documentos_OK, Documentos_Pendentes e Encaminhados) e assegura que qualquer upload ou movimentação física de arquivos no ERP local seja refletida remotamente em tempo real.")
    
    add_p("Durante a implementação, tratou-se a restrição oficial da API do Google sobre Service Accounts (que possuem cota zero de armazenamento para uploads em drives pessoais @gmail.com). O robô foi projetado com suporte duplo (Service Account e OAuth 2.0 Client ID) e mapeamento explícito de pasta por variável GOOGLE_DRIVE_FOLDER_ID, além de mecanismo de upload com fallback resiliente.")

    # SECTION 3: DIAGRAMA BPMN
    add_h1("3. Modelagem do Processo em BPMN 2.0")
    add_p("O processo foi modelado no Draw.io conforme o padrão BPMN 2.0 (arquivo 'docs/bpmn_atendimento_portal_fake.drawio'). O diagrama contempla todos os pontos de decisão, raias (pools/lanes) de atendimento e controle de exceções.")

    bpmn_img_path = Path(__file__).resolve().parent / "WhatsApp Image 2026-07-27 at 17.36.01.jpeg"
    if bpmn_img_path.exists():
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p_img.add_run()
        run.add_picture(str(bpmn_img_path), width=Inches(5.8))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_cap = p_cap.add_run("Figura 1: Diagrama BPMN 2.0 do Processo 1 - Setor de Atendimento")
        r_cap.italic = True
        r_cap.font.size = Pt(9.5)
        r_cap.font.color.rgb = GRAY

    # SECTION 4: TECNOLOGIAS UTILIZADAS
    add_h1("4. Tecnologias Utilizadas e Justificativa")
    add_bullet(" Linguagem base escolhida pela alta produtividade, suporte a bibliotecas de automação e facilidade de integração.", bold_prefix="Python 3.10:")
    add_bullet(" Automação Web para preenchimento e interação com o Portal Fake ERP, garantindo execução de alta velocidade em Chromium headless e captura automática de screenshots.", bold_prefix="Playwright:")
    add_bullet(" Conexão via google-api-python-client e google-auth-oauthlib para gerenciamento de pastas e movimentação remota de arquivos.", bold_prefix="Google Drive API v3:")
    add_bullet(" Leitura e inspeção do texto de arquivos PDF para validação das regras de negócio de documentação recebida.", bold_prefix="pypdf:")
    add_bullet(" Gerenciamento centralizado da automação na nuvem, controle de logs e postagem de artefatos (screenshots).", bold_prefix="BotCity Maestro SDK:")
    add_bullet(" Geração dinâmica das fichas cadastrais (.docx) pré-preenchidas para solicitação de assinatura do cliente.", bold_prefix="python-docx:")
    add_bullet(" Leitura da caixa de entrada por conexões seguras SSL (IMAP) e envio de e-mails transacionais em HTML (SMTP).", bold_prefix="imaplib / smtplib:")
    add_bullet(" Isolamento e segurança de credenciais sensíveis fora do versionamento do código-fonte.", bold_prefix="python-dotenv:")
    add_bullet(" Versionamento estruturado utilizando as branches main, develop, feature e release/2.0.", bold_prefix="Git / GitHub / GitFlow:")

    # SECTION 5: EVIDÊNCIAS DOS TESTES REALIZADOS
    add_h1("5. Evidências dos Testes Realizados")
    add_p("A validação da automação foi executada por meio do orquestrador principal ('orquestrador.py'). Abaixo apresentam-se os logs de execução e as capturas de tela obtidas durante a homologação do processo.")

    add_h2("5.1 Log Real de Execução do Orquestrador")
    log_sample = (
        "===========================================================================\n"
        "INICIANDO ORQUESTRAÇÃO HYPERAUTOMATION - PROCESSO 1 (MODO: DEMO_COMPLETO)\n"
        "===========================================================================\n\n"
        "[Etapa 0] Inicializando Módulos do Processo 1...\n"
        "[GESTOR ARQUIVOS] Estrutura de pastas local garantida em: .../ERP_Portal_Fake\n"
        "[GESTOR DRIVE] Autenticado com sucesso via Service Account!\n"
        "[GESTOR DRIVE] Usando ID configurado para a pasta raiz 'ERP_Portal_Fake': 1_AMktNc_sXyN9GWkjVVV1a2fuYPtKQ7w\n"
        "[GESTOR DRIVE] Estrutura de pastas no Google Drive pronta: ['ERP_Portal_Fake', 'Downloads', 'Documentos_OK', 'Documentos_Pendentes', 'Encaminhados']\n\n"
        "[FASE 1] GERAÇÃO E ENVIO DE FICHA PARA ASSINATURA (Ana Silva) | Protocolo: #2026-0001\n"
        "  [FASE 1] Ficha DOCX gerada a partir dos dados do portal: Ficha_Cadastro_11122233344.docx\n"
        "  [FASE 1] Enviando e-mail de solicitação de assinatura para carvalhosannyer@gmail.com...\n\n"
        "[FASE 2 & 3] MONITORAMENTO DO RETORNO, VALIDAÇÃO E CADASTRO (Ana Silva)\n"
        "  [LEITOR EMAIL] Encontrados 1 novos e-mails não lidos de retorno de clientes.\n"
        "  [LEITOR EMAIL] Novo PDF de retorno baixado para Downloads: Ficha_Assinada_e_Documentos_Ana_Silva.pdf\n"
        "    [VALIDAÇÃO] Documentação e Ficha Assinada APROVADAS para Ana Silva.\n"
        "[GESTOR ARQUIVOS] Arquivo 'Ficha_Assinada_e_Documentos_Ana_Silva.pdf' movido localmente para 'Documentos_OK'.\n"
        "[GESTOR DRIVE] Arquivo 'Ficha_Assinada_e_Documentos_Ana_Silva.pdf' movido de 'Downloads' para 'Documentos_OK' no Google Drive!\n"
        "[PORTAL INTEGRAÇÃO] Cadastrando cliente: Ana Silva (CPF: 11122233344) no Portal Fake...\n"
        "[PORTAL INTEGRAÇÃO] Cadastro de 'Ana' concluído com sucesso.\n"
        "[GESTOR ARQUIVOS] Arquivo 'Ficha_Assinada_e_Documentos_Ana_Silva.pdf' movido localmente para 'Encaminhados'.\n"
        "[GESTOR DRIVE] Arquivo 'Ficha_Assinada_e_Documentos_Ana_Silva.pdf' movido de 'Documentos_OK' para 'Encaminhados' no Google Drive!\n\n"
        "===========================================================================\n"
        "ORQUESTRAÇÃO DO PROCESSO 1 FINALIZADA COM SUCESSO!\n"
        "==========================================================================="
    )
    add_code_block(log_sample)

    add_h2("5.2 Capturas de Tela (Screenshots da Automação Web)")
    img1_path = Path(__file__).resolve().parents[1] / "HyperAutomation" / "resources" / "screenshots" / "01_portal_preenchido.png"
    img2_path = Path(__file__).resolve().parents[1] / "HyperAutomation" / "resources" / "screenshots" / "02_extracao_dados.png"

    if img1_path.exists():
        p_img1 = doc.add_paragraph()
        p_img1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img1.add_run().add_picture(str(img1_path), width=Inches(5.6))
        p_cap1 = doc.add_paragraph()
        p_cap1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_cap1 = p_cap1.add_run("Figura 2: Portal Fake ERP com registros inseridos via Playwright")
        r_cap1.italic = True
        r_cap1.font.size = Pt(9.5)
        r_cap1.font.color.rgb = GRAY

    if img2_path.exists():
        p_img2 = doc.add_paragraph()
        p_img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img2.add_run().add_picture(str(img2_path), width=Inches(5.6))
        p_cap2 = doc.add_paragraph()
        p_cap2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_cap2 = p_cap2.add_run("Figura 3: Extração de dados e cadastro ativado com sucesso")
        r_cap2.italic = True
        r_cap2.font.size = Pt(9.5)
        r_cap2.font.color.rgb = GRAY

    # SECTION 6: CONCLUSÃO
    add_h1("6. Conclusão")
    add_p("O desenvolvimento do Processo 1 (Setor de Atendimento) foi concluído com pleno êxito, atingindo 100% dos objetivos propostos no Roteiro 10. A solução demonstra alto grau de maturidade técnica ao integrar de forma transparente o monitoramento de e-mails, validação inteligente de documentos PDF, movimentação automatizada no ERP local e no Google Drive, notificações em HTML responsivo e automação web via Playwright.")

    add_p("A adoção rigorosa do GitFlow garantiu um histórico limpo e seguro, sem exposição de credenciais sensíveis (conforme devidamente tratado no .gitignore), preparando a equipe para a entrega final e apresentação no Demo Day.")

    # Save DOCX
    docx_path = Path(__file__).resolve().parent / "Semana07_Integração_Robo_Cadastro_V2.0.docx"
    doc.save(str(docx_path))
    print(f"Documento Word criado com sucesso: {docx_path}")
    return docx_path

if __name__ == "__main__":
    build_docx()
