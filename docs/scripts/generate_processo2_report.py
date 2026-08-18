import os
import sys
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def build_processo2_docx():
    doc = docx.Document()

    # Margins (2.5 cm ~ 0.98 in)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.98)
        section.bottom_margin = Inches(0.98)
        section.left_margin = Inches(0.98)
        section.right_margin = Inches(0.98)

    # Styles & Fonts
    style_normal = doc.styles['Normal']
    font_normal = style_normal.font
    font_normal.name = 'Calibri'
    font_normal.size = Pt(11)
    font_normal.color.rgb = RGBColor(0x22, 0x22, 0x22)

    # Palette
    NAVY = RGBColor(0x1B, 0x36, 0x5D)
    DARK_BLUE = RGBColor(0x0F, 0x4C, 0x81)
    GRAY = RGBColor(0x55, 0x55, 0x55)

    def add_title(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(20)
        run.font.color.rgb = NAVY
        p.paragraph_format.space_after = Pt(4)
        return p

    def add_subtitle(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.font.size = Pt(13)
        run.font.color.rgb = GRAY
        run.italic = True
        p.paragraph_format.space_after = Pt(18)
        return p

    def add_h1(text):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(14)
        run.font.color.rgb = NAVY
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(12)
        run.font.color.rgb = DARK_BLUE
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        return p

    def add_p(text, bold_prefix=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.bold = True
            r_bold.font.color.rgb = NAVY
        p.add_run(text)
        return p

    def add_bullet(text, bold_prefix=None):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.bold = True
            r_bold.font.color.rgb = NAVY
        p.add_run(text)
        return p

    def add_code_block(code_text):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F4F6F8"/>')
        tcPr.append(shd)
        
        borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="single" w:sz="4" w:space="0" w:color="D0D7DE"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="D0D7DE"/><w:left w:val="single" w:sz="12" w:space="0" w:color="0F4C81"/><w:right w:val="single" w:sz="4" w:space="0" w:color="D0D7DE"/></w:tcBorders>')
        tcPr.append(borders)

        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(code_text)
        run.font.name = 'Courier New'
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # HEADER & IDENTIFICATION
    add_title("Relatório Técnico: Automação do Processo 2")
    add_subtitle("Organização de Dados | Empresa Portal Fake")

    add_p("Moisés Levy", bold_prefix="Professor: ")
    add_p("Técnicas de Hyperautomation", bold_prefix="Disciplina: ")
    add_p("Portal Fake", bold_prefix="Empresa: ")
    add_p("2 – Organização de Dados", bold_prefix="Processo: ")
    add_p("04 de Agosto de 2026", bold_prefix="Data: ")

    # SECTION 1: IDENTIFICAÇÃO DA EQUIPE
    add_h1("1. IDENTIFICAÇÃO DA EQUIPE")
    add_p("Equipe 1", bold_prefix="Nome da Equipe: ")
    add_p("Integrantes:", bold_prefix="Membros do Grupo: ")
    add_bullet(" Sannyer Cardoso Carvalho Nery (DevOps, Arquitetura Python, GitFlow, Integração Drive/ERP)")
    add_bullet(" Eric Luna Costa (Analista de Processos BPMN & Regras de Negócio)")
    add_bullet(" Daniele Greice Albuquerque e Silva (Desenvolvimento de Automação & Testes)")
    add_bullet(" Kauã Sales Viana (Desenvolvimento de Automação & Suporte de Integração)")

    # SECTION 2: DESCRIÇÃO DA SOLUÇÃO DESENVOLVIDA
    add_h1("2. DESCRIÇÃO DA SOLUÇÃO DESENVOLVIDA")
    add_p("Este projeto contempla o desenvolvimento do segundo módulo da solução de Hyperautomation para a Empresa Portal Fake: a automação do processo de Organização de Dados.")
    add_p("O objetivo central é automatizar de forma ponta a ponta o fluxo de trabalho do setor de Organização de Dados, que consiste em:")
    add_bullet(" Identificar e capturar automaticamente os documentos em PDF aprovados oriundos do Processo 1 localizados na pasta 'Documentos_OK'.", bold_prefix="1. Recepção de Documentos: ")
    add_bullet(" Analisar e extrair programaticamente os dados cadastrais (Nome, Sobrenome e CPF com 11 dígitos) contidos na Ficha de Cadastro em PDF utilizando expressões regulares resilientes (módulo ExtratorPDF).", bold_prefix="2. Extração de Informações: ")
    add_bullet(" Consolidar e gravar as informações extraídas na Planilha Mestra (Planilha_Mestra.xlsx), realizando validação de duplicidades de CPF e registrando timestamp de processamento.", bold_prefix="3. Consolidação e Tratamento de Duplicidades: ")
    add_bullet(" Arquivar com segurança os documentos originais processados, movendo-os localmente para a pasta 'Arquivados' e replicando instantaneamente a alteração no Google Drive via Google Drive API v3.", bold_prefix="4. Arquivamento Físico e em Nuvem: ")

    add_p("A solução foi desenvolvida utilizando a linguagem Python, seguindo os princípios de Arquitetura Limpa, desacoplamento de responsabilidades e controle de versionamento rigoroso com GitFlow.")

    # SECTION 3: DIAGRAMA BPMN
    add_h1("3. DIAGRAMA BPMN")
    add_p("O fluxo do Processo 2 foi totalmente mapeado segundo a notação BPMN 2.0 (arquivo 'docs/bpmn_atendimento_portal_fake2.drawio'). Ele detalha as etapas de gatilho, leitura de documentos aprovados, extração de texto PDF, validação de unicidade de CPF na planilha e o arquivamento final.")

    bpmn2_img = Path(__file__).resolve().parents[1] / "img" / "bpmn_atendimento_portal_fake2.jpg"
    if bpmn2_img.exists():
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.add_run().add_picture(str(bpmn2_img), width=Inches(5.8))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_cap = p_cap.add_run("Figura 1: Diagrama BPMN 2.0 do Processo 2 - Organização de Dados")
        r_cap.italic = True
        r_cap.font.size = Pt(9.5)
        r_cap.font.color.rgb = GRAY

    # SECTION 4: ARQUITETURA DA SOLUÇÃO
    add_h1("4. ARQUITETURA DA SOLUÇÃO")
    add_p("A arquitetura do Processo 2 é organizada em módulos especializados com alta coesão:")
    add_bullet(" Script principal responsável pela orquestração de todo o fluxo, conectando os módulos de gestão de arquivos, extração de texto e atualização de planilha.", bold_prefix="orquestrador_processo2.py: ")
    add_bullet(" Responsável por abrir os arquivos PDF de 'Documentos_OK', extrair o texto via 'pypdf' e aplicar padrões Regex para extrair Nome, Sobrenome e CPF sanitizado (11 dígitos).", bold_prefix="processo_organizacao/extrator_dados.py (ExtratorPDF): ")
    add_bullet(" Responsável por inicializar e manipular o arquivo 'Planilha_Mestra.xlsx' utilizando Pandas e OpenPyXL. Valida a existência prévia de registros pelo CPF para evitar entradas duplicadas.", bold_prefix="processo_organizacao/planilha_mestra.py (GerenciadorPlanilha): ")
    add_bullet(" Responsável por garantir a criação das pastas do ERP e gerenciar a movimentação física dos arquivos locais e sincronização no Google Drive.", bold_prefix="processo_atendimento/gestor_arquivos.py (GestorArquivos): ")
    add_bullet(" Interface com a API do Google Drive (v3), efetuando autenticação OAuth 2.0 / Service Account e movimentação remota dos arquivos para a pasta 'Arquivados'.", bold_prefix="processo_atendimento/gestor_drive.py (GestorDrive): ")

    # SECTION 5: TECNOLOGIAS UTILIZADAS
    add_h1("5. TECNOLOGIAS UTILIZADAS")
    
    table_tech = doc.add_table(rows=6, cols=2)
    table_tech.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_tech.autofit = False

    hdr_t = table_tech.rows[0].cells
    hdr_t[0].text = "Tecnologia / Biblioteca"
    hdr_t[1].text = "Aplicação e Justificativa no Processo 2"
    for cell in hdr_t:
        tcPr = cell._tc.get_or_add_tcPr()
        tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="1B365D"/>'))
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    tech_data = [
        ("Python 3.10", "Linguagem core do projeto, escolhida pela riqueza de ecossistema para automação e manipulação de dados."),
        ("pypdf (v4.0+)", "Extração eficiente de texto bruto a partir de arquivos PDF para parsing de formulários cadastrais."),
        ("Pandas & OpenPyXL", "Criação, leitura e atualização da Planilha_Mestra.xlsx com alta performance e tipagem rigorosa de dados (preservação de zeros do CPF)."),
        ("Google Drive API v3 / PyDrive2", "Sincronização em tempo real da estrutura de pastas físicas do ERP com a nuvem."),
        ("Git / GitFlow", "Versionamento estruturado com branches organizadas (main, develop, feature/input_data).")
    ]

    for idx, (t_name, t_desc) in enumerate(tech_data, start=1):
        r_cells = table_tech.rows[idx].cells
        r_cells[0].text = t_name
        r_cells[1].text = t_desc
        r_cells[0].paragraphs[0].runs[0].bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # SECTION 6: ESTRUTURA DO PROJETO
    add_h1("6. ESTRUTURA DO PROJETO")
    add_p("A estrutura de diretórios do projeto segue o padrão modular estabelecido:")

    tree_str = (
        "hyperautomation-equipe-1/\n"
        "├── ERP_Portal_Fake/\n"
        "│   ├── Downloads/\n"
        "│   ├── Documentos_OK/\n"
        "│   ├── Documentos_Pendentes/\n"
        "│   ├── Sistema_Integrador_Portal_Fake/\n"
        "│   │   └── Planilha_Mestra.xlsx\n"
        "│   └── Arquivados/\n"
        "├── HyperAutomation/\n"
        "│   ├── requirements.txt\n"
        "│   └── source/\n"
        "│       ├── orquestrador.py (Processo 1)\n"
        "│       ├── orquestrador_processo2.py (Processo 2)\n"
        "│       ├── processo_atendimento/\n"
        "│       │   ├── gestor_arquivos.py\n"
        "│       │   └── gestor_drive.py\n"
        "│       └── processo_organizacao/\n"
        "│           ├── extrator_dados.py\n"
        "│           └── planilha_mestra.py\n"
        "└── docs/\n"
        "    ├── bpmn_atendimento_portal_fake2.drawio\n"
        "    └── Relatorio_Tecnico_Automacao_Processo_2.docx"
    )
    add_code_block(tree_str)

    # SECTION 7: EVIDÊNCIAS DOS TESTES REALIZADOS
    add_h1("7. EVIDÊNCIAS DOS TESTES REALIZADOS")
    add_p("A execução do Processo 2 foi validada com sucesso através de testes automatizados e simulações com documentos reais do portal.")

    add_h2("7.1 Log Real de Execução da Orquestração do Processo 2")
    log_p2 = (
        "===========================================================================\n"
        "INICIANDO ORQUESTRAÇÃO HYPERAUTOMATION - PROCESSO 2 (Organização de Dados)\n"
        "===========================================================================\n"
        "[GESTOR DRIVE] Autenticado com sucesso via OAuth 2.0 Client ID!\n"
        "[GESTOR DRIVE] Conexão com Google Drive API v3 estabelecida com sucesso!\n"
        "[GESTOR ARQUIVOS] Estrutura de pastas local garantida em: .../ERP_Portal_Fake\n"
        "[GESTOR DRIVE] Usando ID configurado para a pasta raiz 'ERP_Portal_Fake': 1_AMktNc_sXyN9GWkjVVV1a2fuYPtKQ7w\n"
        "[GESTOR DRIVE] Estrutura de pastas no Google Drive pronta: ['ERP_Portal_Fake', 'Downloads', 'Documentos_OK', 'Documentos_Pendentes', 'Arquivados']\n"
        "[PROCESSO 2] Encontrados 2 documentos aprovados. Iniciando processamento...\n\n"
        "Processando arquivo: ficha_documento_comprovante_assinado (1).pdf\n"
        "  -> Dados Extraídos: Nome: Sannyer Nery | CPF: 03536054250\n"
        "[PLANILHA MESTRA] Planilha criada em: .../Sistema_Integrador_Portal_Fake/Planilha_Mestra.xlsx\n"
        "[PLANILHA MESTRA] Registro de Sannyer adicionado com sucesso.\n"
        "[GESTOR ARQUIVOS] Arquivo 'ficha_documento_comprovante_assinado (1).pdf' arquivado com sucesso.\n"
        "[GESTOR DRIVE] Arquivo 'ficha_documento_comprovante_assinado (1).pdf' movido de 'Documentos_OK' para 'Arquivados' no Google Drive!\n\n"
        "Processando arquivo: ficha_documento_comprovante_assinado.pdf\n"
        "  -> Dados Extraídos: Nome: Sannyer Nery | CPF: 03536054250\n"
        "[PLANILHA MESTRA] Aviso: CPF 03536054250 já cadastrado. Ignorando duplicidade.\n"
        "[GESTOR ARQUIVOS] Arquivo 'ficha_documento_comprovante_assinado.pdf' arquivado com sucesso.\n"
        "[GESTOR DRIVE] Arquivo 'ficha_documento_comprovante_assinado.pdf' movido de 'Documentos_OK' para 'Arquivados' no Google Drive!\n\n"
        "===========================================================================\n"
        "ORQUESTRAÇÃO DO PROCESSO 2 FINALIZADA COM SUCESSO!\n"
        "==========================================================================="
    )
    add_code_block(log_p2)

    add_h2("7.2 Espaço Reservado para Evidências Complementares")
    add_p("Abaixo você pode inserir capturas de tela adicionais (Screenshots da Planilha_Mestra.xlsx aberta no Excel, telas do Google Drive com a pasta Arquivados, etc.):")

    # Table for placeholder screenshots
    table_evid = doc.add_table(rows=2, cols=2)
    table_evid.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_evid.autofit = False

    h_evid = table_evid.rows[0].cells
    h_evid[0].text = "Item / Evidência"
    h_evid[1].text = "Espaço para Imagem / Captura de Tela"
    for cell in h_evid:
        tcPr = cell._tc.get_or_add_tcPr()
        tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="0F4C81"/>'))
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    row1 = table_evid.rows[1].cells
    row1[0].text = "Planilha_Mestra.xlsx\n(Registros Consolidados no Excel)"
    row1[1].text = "[INSERIR AQUI A IMAGEM DA PLANILHA MESTRA NO EXCEL]"

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # SECTION 8: ORGANIZAÇÃO NO GOOGLE DRIVE
    add_h1("8. ORGANIZAÇÃO NO GOOGLE DRIVE")
    add_p("A integração automatizada com a nuvem garante que a pasta raiz 'ERP_Portal_Fake' no Google Drive reflita fielmente a estrutura física local de pastas do ERP:")
    add_bullet(" Downloads/ : Arquivos baixados de e-mails em triagem.", bold_prefix="• ")
    add_bullet(" Documentos_OK/ : Documentos aprovados no Processo 1 aguardando processamento do Processo 2.", bold_prefix="• ")
    add_bullet(" Documentos_Pendentes/ : Documentos reprovados ou com pendências documentais.", bold_prefix="• ")
    add_bullet(" Arquivados/ : Documentos cujos dados já foram extraídos e gravados na Planilha Mestra pelo Processo 2.", bold_prefix="• ")

    # SECTION 9: CONCLUSÃO
    add_h1("9. CONCLUSÃO")
    add_p("A equipe concluiu com êxito o desenvolvimento do segundo módulo de automação para a Empresa Portal Fake, entregando uma solução robusta, escalável e alinhada aos objetivos de Hyperautomation.")
    add_p("O sistema automatiza integralmente o fluxo de Organização de Dados, desde o monitoramento da pasta de documentos aprovados até o arquivamento seguro dos registros, destacando-se pela precisão na extração de informações, validação rigorosa dos dados e integração fluida com a Planilha Mestra.")
    add_p("A arquitetura modular adotada garante manutenibilidade e escalabilidade, enquanto o uso do GitFlow proporcionou um desenvolvimento organizado e com histórico claro de alterações. Os testes realizados asseguram a confiabilidade do sistema em diferentes cenários, reduzindo significativamente os riscos de falhas em produção.")
    add_p("Dessa forma, a solução desenvolvida não apenas cumpre os objetivos propostos, mas também estabelece uma base sólida para a evolução contínua do ecossistema de automação da Empresa Portal Fake, demonstrando o valor da Hyperautomation na otimização de processos corporativos.")

    # Save
    out_docx = Path(__file__).resolve().parents[1] / "Relatorio_Tecnico_Automacao_Processo_2.docx"
    doc.save(str(out_docx))
    print(f"[SUCESSO] Relatório DOCX do Processo 2 gerado em: {out_docx}")
    return out_docx

if __name__ == "__main__":
    build_processo2_docx()
