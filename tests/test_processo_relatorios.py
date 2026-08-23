"""
Testes Automatizados - Processo 5 (Relatórios e Gerência)
Validação de consolidação de dados, métricas executivas, relatórios Excel/MD/JSON e versionamento.
"""
import sys
import json
import pytest
from pathlib import Path
from unittest.mock import MagicMock, patch
from openpyxl import Workbook, load_workbook

# Adiciona diretórios ao path
BASE_DIR = Path(__file__).resolve().parent.parent / "HyperAutomation" / "source"
sys.path.append(str(BASE_DIR))
sys.path.append(str(BASE_DIR.parent))

from processo_relatorios.consolidador import ConsolidadorDados
from processo_relatorios.metricas import CalculadoraMetricas
from processo_relatorios.gerador_relatorio import GeradorRelatorios
from processo_relatorios.logger_relatorios import LoggerRelatorios
from processo_relatorios.orquestrador_relatorios import ProcessoRelatorios, executar_processo5


@pytest.fixture
def planilhas_teste(tmp_path):
    """Cria planilhas temporárias de teste para a Planilha Mestra e SAC."""
    p_mestra = tmp_path / "Planilha_Mestra.xlsx"
    p_sac = tmp_path / "atendimentos_sac.xlsx"

    # 1. Planilha Mestra
    wb_m = Workbook()
    ws_m = wb_m.active
    ws_m.title = "Planilha Mestra"
    ws_m.append([
        "CPF", "Nome", "Data de Nascimento", "Endereço", "E-mail",
        "Telefone", "Status", "Data de Processamento", "Observações", "Protocolo"
    ])
    ws_m.append([
        "10000012482", "Ana Silva", "1995-01-01", "Manaus/AM", "ana@email.com",
        "(92) 99888-1122", "CONCLUIDO_P3", "20/08/2026 10:00", "Cadastrado via RPA", "PROT-001"
    ])
    ws_m.append([
        "10000012756", "Bruno Souza", "1990-05-10", "Manaus/AM", "bruno@email.com",
        "(92) 99777-2233", "DUPLICADO_P3", "20/08/2026 10:05", "CPF já cadastrado", "PROT-002"
    ])
    ws_m.append([
        "10000013441", "Carla Lima", "1988-12-20", "Manaus/AM", "carla@email.com",
        "(92) 99666-3344", "ERRO_VALIDACAO_P3", "20/08/2026 10:10", "CPF inválido", "PROT-003"
    ])
    wb_m.save(p_mestra)

    # 2. Planilha SAC
    wb_s = Workbook()
    ws_s = wb_s.active
    ws_s.title = "Atendimentos"
    ws_s.append([
        "Chave Atendimento", "Protocolo", "CPF", "Nome", "E-mail",
        "Status Cadastro", "Tipo Atendimento", "Status Atendimento", "Data Atendimento", "Observação"
    ])
    ws_s.append([
        "PROT-001_CADASTRO_OK", "PROT-001", "10000012482", "Ana Silva", "ana@email.com",
        "CONCLUIDO_P3", "CADASTRO_OK", "COMUNICADO", "20/08/2026 10:15", "E-mail enviado com sucesso."
    ])
    ws_s.append([
        "PROT-002_CADASTRO_OK", "PROT-002", "10000012756", "Bruno Souza", "bruno@email.com",
        "DUPLICADO_P3", "CADASTRO_OK", "PENDENTE_CONTATO", "", "Sem e-mail cadastrado."
    ])
    wb_s.save(p_sac)

    return {
        "mestra": p_mestra,
        "sac": p_sac,
        "dir_saida": tmp_path / "relatorios_out"
    }


class TestConsolidadorDados:
    """Testes do módulo de carregamento e cruzamento de dados."""

    def test_carregar_planilha_mestra(self, planilhas_teste):
        consolidador = ConsolidadorDados(
            caminho_mestra=planilhas_teste["mestra"],
            caminho_sac=planilhas_teste["sac"]
        )
        dados = consolidador.carregar_planilha_mestra()
        assert len(dados) == 3
        assert dados[0]["Nome"] == "Ana Silva"
        assert dados[0]["cpf_limpo"] == "10000012482"
        assert dados[0]["Status"] == "CONCLUIDO_P3"

    def test_carregar_planilha_sac(self, planilhas_teste):
        consolidador = ConsolidadorDados(
            caminho_mestra=planilhas_teste["mestra"],
            caminho_sac=planilhas_teste["sac"]
        )
        dados = consolidador.carregar_planilha_sac()
        assert len(dados) == 2
        assert dados[0]["Protocolo"] == "PROT-001"
        assert dados[0]["Status Atendimento"] == "COMUNICADO"

    def test_consolidar_cruzamento_completo(self, planilhas_teste):
        consolidador = ConsolidadorDados(
            caminho_mestra=planilhas_teste["mestra"],
            caminho_sac=planilhas_teste["sac"]
        )
        res = consolidador.consolidar()
        assert res["total_clientes"] == 3
        assert res["total_mestra"] == 3
        assert res["total_sac"] == 2

        clientes = res["clientes"]
        ana = next(c for c in clientes if c["cpf_limpo"] == "10000012482")
        assert ana["status_cadastro"] == "CONCLUIDO_P3"
        assert ana["status_sac"] == "COMUNICADO"
        assert "PROCESSO 4" in ana["fase_concluida"]

        bruno = next(c for c in clientes if c["cpf_limpo"] == "10000012756")
        assert bruno["status_cadastro"] == "DUPLICADO_P3"
        assert bruno["status_sac"] == "PENDENTE_CONTATO"

        carla = next(c for c in clientes if c["cpf_limpo"] == "10000013441")
        assert carla["status_sac"] == "NÃO INICIADO"

    def test_consolidar_com_arquivos_inexistentes(self, tmp_path):
        consolidador = ConsolidadorDados(
            caminho_mestra=tmp_path / "inexistente_mestra.xlsx",
            caminho_sac=tmp_path / "inexistente_sac.xlsx"
        )
        res = consolidador.consolidar()
        assert res["total_clientes"] == 0
        assert res["clientes"] == []


class TestCalculadoraMetricas:
    """Testes do motor analítico de KPIs."""

    def test_calculo_metricas(self, planilhas_teste):
        consolidador = ConsolidadorDados(
            caminho_mestra=planilhas_teste["mestra"],
            caminho_sac=planilhas_teste["sac"]
        )
        dados_consolidados = consolidador.consolidar()
        metricas = CalculadoraMetricas.calcular(dados_consolidados)

        # Validação do Resumo
        assert metricas["resumo_geral"]["total_clientes_processados"] == 3
        assert metricas["resumo_geral"]["total_registros_planilha_mestra"] == 3
        assert metricas["resumo_geral"]["total_registros_sac"] == 2

        # Validação de KPIs de Cadastro
        kpi_cad = metricas["kpis_cadastro_processo3"]
        assert kpi_cad["cadastrados_com_sucesso"] == 1
        assert kpi_cad["duplicados_identificados"] == 1
        assert kpi_cad["erros_validacao"] == 1
        assert kpi_cad["taxa_aprovacao_percentual"] == 33.33

        # Validação de KPIs de SAC
        kpi_sac = metricas["kpis_sac_processo4"]
        assert kpi_sac["total_atendimentos_sac"] == 2
        assert kpi_sac["comunicados_com_sucesso"] == 1
        assert kpi_sac["pendentes_contato_fallback"] == 1
        assert kpi_sac["taxa_sucesso_comunicacao_percentual"] == 50.0

        # Validação de Eficiência Global
        kpi_efi = metricas["kpis_eficiencia_global"]
        assert kpi_efi["taxa_conclusao_ponta_a_ponta_percentual"] == 33.33

    def test_calculo_metricas_base_vazia(self):
        metricas = CalculadoraMetricas.calcular({"clientes": []})
        assert metricas["resumo_geral"]["total_clientes_processados"] == 0
        assert metricas["kpis_cadastro_processo3"]["taxa_aprovacao_percentual"] == 0.0
        assert metricas["kpis_sac_processo4"]["taxa_sucesso_comunicacao_percentual"] == 0.0


class TestGeradorRelatorios:
    """Testes de geração de relatórios Excel, Markdown e JSON com versionamento."""

    def test_gerar_todos_relatorios(self, planilhas_teste):
        consolidador = ConsolidadorDados(
            caminho_mestra=planilhas_teste["mestra"],
            caminho_sac=planilhas_teste["sac"]
        )
        dados = consolidador.consolidar()
        metricas = CalculadoraMetricas.calcular(dados)

        gerador = GeradorRelatorios(dir_saida=planilhas_teste["dir_saida"])
        resultado = gerador.gerar_todos(dados, metricas)

        # 1. Verifica arquivos gerados
        assert resultado["excel"].exists()
        assert resultado["markdown"].exists()
        assert resultado["json"].exists()

        # 2. Verifica cópias 'latest'
        dir_out = planilhas_teste["dir_saida"]
        assert (dir_out / "relatorio_gerencial_latest.xlsx").exists()
        assert (dir_out / "relatorio_gerencial_latest.md").exists()
        assert (dir_out / "relatorio_gerencial_latest.json").exists()

        # 3. Valida estrutura do Excel gerado
        wb = load_workbook(resultado["excel"])
        assert "Dashboard Gerencial" in wb.sheetnames
        assert "Base Consolidada" in wb.sheetnames
        assert "Distribuição de Status" in wb.sheetnames

        ws_dash = wb["Dashboard Gerencial"]
        assert "RELATÓRIO GERENCIAL - HYPERAUTOMATION" in str(ws_dash["A1"].value)

        ws_base = wb["Base Consolidada"]
        assert ws_base.max_row == 4  # Cabeçalho + 3 clientes

        # 4. Valida conteúdo do Markdown
        conteudo_md = resultado["markdown"].read_text(encoding="utf-8")
        assert "# 📊 Relatório Gerencial de HyperAutomation" in conteudo_md
        assert "Ana Silva" not in conteudo_md  # Resumo executivo foca em tabelas e KPIs
        assert "Processo 3 — Cadastro" in conteudo_md

        # 5. Valida payload do JSON
        conteudo_json = json.loads(resultado["json"].read_text(encoding="utf-8"))
        assert "metadata" in conteudo_json
        assert "metricas" in conteudo_json
        assert len(conteudo_json["clientes"]) == 3


class TestProcessoRelatoriosOrquestracao:
    """Testes de execução integrada do Processo 5."""

    def test_executar_processo5_completo(self, planilhas_teste):
        resultado = executar_processo5(
            caminho_planilha_mestra=planilhas_teste["mestra"],
            caminho_planilha_sac=planilhas_teste["sac"],
            dir_saida=planilhas_teste["dir_saida"]
        )

        assert resultado["sucesso"] is True
        assert resultado["total_clientes"] == 3
        assert "relatorios" in resultado
        assert resultado["relatorios"]["excel"].exists()

    def test_executar_processo5_com_falha_capturada(self, tmp_path):
        processo = ProcessoRelatorios()
        # Força erro mockando gerador
        processo.gerador.gerar_todos = MagicMock(side_effect=IOError("Disco cheio simulado"))
        res = processo.executar()

        assert res["sucesso"] is False
        assert "Disco cheio simulado" in res["erro"]
